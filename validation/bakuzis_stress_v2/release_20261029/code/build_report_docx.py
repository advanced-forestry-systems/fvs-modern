#!/usr/bin/env python3
"""build_report_docx.py: render the catalog technical report markdown to docx in the CRSF house style.
Usage: python3 build_report_docx.py --md report.md --out report.docx
House style: tables centered, bold centered header, rules above header, below header and below last row only,
units in headers; table captions above, figure captions below, label bold; species italics not needed here.
"""
import argparse, re, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

GREEN = RGBColor(0x1a, 0x3d, 0x28); GOLD = RGBColor(0xc5, 0xa5, 0x5a)

def set_cell_borders(cell, top=None, bottom=None):
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders")
    for side, val in (("top", top), ("bottom", bottom), ("left", "nil"), ("right", "nil"), ("insideH", "nil"), ("insideV", "nil")):
        el = OxmlElement(f"w:{side}")
        if val in (None, "nil"): el.set(qn("w:val"), "nil")
        else: el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(val)); el.set(qn("w:color"), "000000")
        b.append(el)
    tcPr.append(b)

def add_runs(par, text):
    # minimal inline markdown: **bold**, *italic*, `code`
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not tok: continue
        if tok.startswith("**"): r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("`"): r = par.add_run(tok[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
        elif tok.startswith("*"): r = par.add_run(tok[1:-1]); r.italic = True
        else: par.add_run(tok)

def add_table(doc, rows, caption):
    if len(rows) > 18: doc.add_page_break()
    cap = doc.add_paragraph(); add_runs(cap, caption); cap.paragraph_format.space_after = Pt(4); cap.paragraph_format.keep_with_next = True
    hdr, body = rows[0], rows[1:]
    t = doc.add_table(rows=len(rows), cols=len(hdr)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = True
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
            r = p.add_run(val); r.font.size = Pt(8.5); r.bold = (i == 0)
            top = 8 if i == 0 else None; bottom = 6 if i == 0 else (8 if i == len(rows) - 1 else None)
            set_cell_borders(c, top, bottom)
    # column widths proportional to the longest cell, with a floor, summing to the text width
    ncol = len(hdr); L = [max(max(max(len(w) for w in r[j].split()) for r in rows) + 2, min(max(len(r[j]) for r in rows[1:]), 24) * 0.95, 7) for j in range(ncol)]; tot = sum(L); widths = [Cm(16.5 * x / tot) for x in L]
    t.autofit = False
    tblPr = t._tbl.tblPr; lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    for j in range(ncol):
        t.columns[j].width = widths[j]
        for i in range(len(rows)): t.cell(i, j).width = widths[j]
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--md", required=True); ap.add_argument("--out", required=True); a = ap.parse_args()
    base = os.path.dirname(os.path.abspath(a.md))
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11); st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.08
    for s in doc.sections: s.left_margin = s.right_margin = Cm(2.3); s.top_margin = s.bottom_margin = Cm(2.2)
    lines = open(a.md, encoding="utf8").read().split("\n"); i = 0; first_h1 = True
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("# "):
            p = doc.add_paragraph(); r = p.add_run(ln[2:]); r.bold = True; r.font.size = Pt(17); r.font.color.rgb = GREEN; p.paragraph_format.space_after = Pt(8)
            i += 1; continue
        if ln.startswith("## "):
            p = doc.add_paragraph(); r = p.add_run(ln[3:]); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = GREEN
            p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4); p.paragraph_format.keep_with_next = True
            i += 1; continue
        if ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells): rows.append(cells)
                i += 1
            # caption is the previous non-empty paragraph starting with **Table
            cap = pending.pop() if pending else ""
            add_table(doc, rows, cap); continue
        m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", ln)
        if m:
            path = os.path.normpath(os.path.join(base, m.group(2)))
            doc.add_picture(path, width=Cm(13.5)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1; continue
        if ln.strip() == "": i += 1; continue
        if ln.startswith("**Table"):
            pending.append(ln); i += 1; continue
        p = doc.add_paragraph(); add_runs(p, ln)
        if ln.startswith("**Figure"): p.paragraph_format.space_after = Pt(10)
        if first_h1 and ln.startswith("Aaron R. Weiskittel"):
            for r in p.runs: r.font.size = Pt(10.5); r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            first_h1 = False
        if ln.startswith(("Crookston", "Leary", "Reineke")): p.paragraph_format.left_indent = Cm(1); p.paragraph_format.first_line_indent = Cm(-1);
        i += 1
    doc.save(a.out); print("wrote", a.out)

pending = []
if __name__ == "__main__":
    main()
